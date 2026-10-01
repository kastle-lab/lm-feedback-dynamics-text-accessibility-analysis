# Compare LLM judge scores with combined human ratings on overlap rows.
# For each rewritten text with one or more human ratings, an LLM score is counted as correct when it matches any human score for the same rubric. This makes rows with human disagreement permissive: if humans gave 3 and 4, either 3 or 4 is treated as agreement.

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_RUBRICS = {
    "Fidelity": "fidelity_score",
    "Coverage": "coverage_score",
    "Structural Coherence": "structural_coherence_score",
    "Readability": "readability_score",
}


@dataclass
class HumanItem:
    match_text_key: str
    ratings: dict[str, list[int]]
    rows: list[dict[str, object]]


def normalize_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def display_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def normalize_text(value: object) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def text_key(value: object) -> str:
    return hashlib.sha1(normalize_text(value).encode("utf-8")).hexdigest()


def stable_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value).strip())


def human_item_key(metadata: dict[str, object], rewritten_text: object) -> str:
    parts = [
        f"{name}={stable_value(metadata.get(name))}"
        for name in ["chapter_id", "chapter_key", "model", "method", "iteration"]
    ]
    parts.append(f"rewritten_text={normalize_text(rewritten_text)}")
    return hashlib.sha1("||".join(parts).encode("utf-8")).hexdigest()


def parse_rating(value: object) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        rating = value
    elif isinstance(value, float) and value.is_integer():
        rating = int(value)
    else:
        text = str(value).strip()
        if not re.fullmatch(r"[1-5](?:\.0+)?", text):
            return None
        rating = int(float(text))
    return rating if 1 <= rating <= 5 else None


def find_column(headers: list[object], wanted: str) -> int | None:
    target = normalize_header(wanted)
    for idx, header in enumerate(headers):
        if normalize_header(header) == target:
            return idx
    return None


def model_name_from_path(path: Path) -> str:
    stem = path.stem
    if stem.startswith("judge_results_"):
        return stem[len("judge_results_") :]
    return stem


def read_human_items(
    input_dir: Path,
    glob_pattern: str,
    sheet_name: str,
    rewritten_column: str,
    rubrics: dict[str, str],
) -> tuple[dict[str, HumanItem], list[list[object]]]:
    items: dict[str, HumanItem] = {}
    issues: list[list[object]] = []
    files = sorted(input_dir.glob(glob_pattern))
    if not files:
        raise FileNotFoundError(f"No human annotation files matched {input_dir / glob_pattern}")

    for path in files:
        workbook = load_workbook(path, data_only=True, read_only=True)
        if sheet_name not in workbook.sheetnames:
            issues.append([path.name, "", "", "missing_sheet", f"Missing sheet: {sheet_name}"])
            continue

        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration:
            issues.append([path.name, "", "", "empty_sheet", "Sheet has no rows"])
            continue

        rewritten_idx = find_column(headers, rewritten_column)
        if rewritten_idx is None:
            issues.append([path.name, "", "", "missing_rewritten_column", rewritten_column])
            continue

        rubric_indices = {rubric: find_column(headers, rubric) for rubric in rubrics}
        missing = [rubric for rubric, idx in rubric_indices.items() if idx is None]
        if missing:
            issues.append([path.name, "", "", "missing_rubric_columns", ", ".join(missing)])
            continue

        metadata_indices = {
            name: find_column(headers, name)
            for name in ["chapter_id", "chapter_key", "model", "method", "iteration"]
        }

        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            rewritten_text = row[rewritten_idx] if rewritten_idx < len(row) else ""
            match_key = text_key(rewritten_text)
            metadata = {
                name: row[idx] if idx is not None and idx < len(row) else None
                for name, idx in metadata_indices.items()
            }
            key = human_item_key(metadata, rewritten_text)
            if not match_key:
                continue
            item = items.setdefault(
                key,
                HumanItem(match_text_key=match_key, ratings={rubric: [] for rubric in rubrics}, rows=[]),
            )
            item.rows.append(
                {
                    "source_file": path.name,
                    "row_number": row_number,
                    **metadata,
                }
            )
            for rubric, idx in rubric_indices.items():
                raw_value = row[idx] if idx is not None and idx < len(row) else None
                rating = parse_rating(raw_value)
                if rating is None:
                    issues.append([path.name, "", row_number, "invalid_rating", f"{rubric}: {raw_value!r}"])
                else:
                    item.ratings[rubric].append(rating)

    return items, issues


def parse_judgment(value: object, rubrics: dict[str, str]) -> dict[str, int] | None:
    if isinstance(value, dict):
        parsed = value
    else:
        text = str(value or "").strip()
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError:
            match = re.search(r"\{.*\}", text, flags=re.S)
            if not match:
                return None
            try:
                parsed = json.loads(match.group(0))
            except json.JSONDecodeError:
                return None

    scores: dict[str, int] = {}
    for rubric, judge_field in rubrics.items():
        rating = parse_rating(parsed.get(judge_field))
        if rating is None:
            return None
        scores[rubric] = rating
    return scores


def modal_prediction(
    predictions: list[dict[str, int | None]], rubrics: dict[str, str]
) -> tuple[dict[str, int | None], int]:
    collapsed = {}
    conflicts = 0
    for rubric in rubrics:
        values = [prediction[rubric] for prediction in predictions]
        if len(set(values)) > 1:
            conflicts += 1
        counts = Counter(values)
        max_count = max(counts.values())
        modes = {value for value, count in counts.items() if count == max_count}
        collapsed[rubric] = next(value for value in values if value in modes)
    return collapsed, conflicts


def read_llm_predictions(
    judge_dir: Path,
    glob_pattern: str,
    human_items: dict[str, HumanItem],
    rubrics: dict[str, str],
) -> tuple[dict[str, dict[str, dict[str, int | None]]], list[list[object]], dict[str, dict[str, int]]]:
    predictions_by_model_text: dict[str, dict[str, list[dict[str, int | None]]]] = defaultdict(lambda: defaultdict(list))
    validation_rows: list[list[object]] = []
    duplicate_stats: dict[str, dict[str, int]] = {}
    human_keys_by_text: dict[str, list[str]] = defaultdict(list)
    for human_key, human_item in human_items.items():
        human_keys_by_text[human_item.match_text_key].append(human_key)
    files = sorted(judge_dir.glob(glob_pattern))
    if not files:
        raise FileNotFoundError(f"No judge JSONL files matched {judge_dir / glob_pattern}")

    for path in files:
        model = model_name_from_path(path)
        with path.open(encoding="utf-8") as handle:
            for line_number, line in enumerate(handle, start=1):
                try:
                    record = json.loads(line)
                except json.JSONDecodeError as exc:
                    validation_rows.append([path.name, model, line_number, "invalid_json", str(exc)])
                    continue
                if record.get("status") != "ok":
                    validation_rows.append([path.name, model, line_number, "non_ok_status", record.get("status")])
                    continue
                source_text = (record.get("source_row") or {}).get("text", "")
                key = text_key(source_text)
                if key not in human_keys_by_text:
                    continue
                scores = parse_judgment(record.get("judgment"), rubrics)
                if scores is None:
                    validation_rows.append([path.name, model, line_number, "invalid_judgment", "Could not parse scores"])
                    scores = {rubric: None for rubric in rubrics}
                predictions_by_model_text[model][key].append(scores)

    collapsed_by_model: dict[str, dict[str, dict[str, int | None]]] = {}
    for model, predictions_by_text in predictions_by_model_text.items():
        collapsed_by_model[model] = {}
        duplicate_keys = 0
        conflicting_duplicate_metric_predictions = 0
        for text_match_key, predictions in predictions_by_text.items():
            if len(predictions) > 1:
                duplicate_keys += 1
            collapsed, conflicts = modal_prediction(predictions, rubrics)
            conflicting_duplicate_metric_predictions += conflicts
            for human_key in human_keys_by_text[text_match_key]:
                collapsed_by_model[model][human_key] = collapsed
        duplicate_stats[model] = {
            "duplicate_overlap_items": duplicate_keys,
            "conflicting_duplicate_metric_predictions": conflicting_duplicate_metric_predictions,
        }

    return collapsed_by_model, validation_rows, duplicate_stats


def pct(numerator: int | float, denominator: int | float) -> float | None:
    return numerator / denominator if denominator else None


def weighted_kappa(pairs: list[tuple[int, int]], min_rating: int = 1, max_rating: int = 5) -> float | None:
    if not pairs:
        return None
    categories = list(range(min_rating, max_rating + 1))
    n = len(pairs)
    width = max_rating - min_rating
    left_counts = Counter(left for left, _ in pairs)
    right_counts = Counter(right for _, right in pairs)
    observed = 0.0
    expected = 0.0
    for i in categories:
        for j in categories:
            weight = 0.0 if width == 0 else ((i - j) / width) ** 2
            observed_count = sum(1 for left, right in pairs if left == i and right == j)
            expected_count = left_counts[i] * right_counts[j] / n
            observed += weight * observed_count / n
            expected += weight * expected_count / n
    if math.isclose(expected, 0.0):
        return 1.0 if math.isclose(observed, 0.0) else None
    return 1.0 - observed / expected


def nearest_human_score(prediction: int, allowed: set[int]) -> int:
    return min(allowed, key=lambda value: (abs(prediction - value), value))


def summarize(
    human_items: dict[str, HumanItem],
    predictions_by_model: dict[str, dict[str, dict[str, int | None]]],
    duplicate_stats: dict[str, dict[str, int]],
    rubrics: dict[str, str],
) -> tuple[
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
]:
    summary_rows: list[list[object]] = []
    metric_rows: list[list[object]] = []
    distribution_rows: list[list[object]] = []
    bias_rows: list[list[object]] = []
    group_rows: list[list[object]] = []
    method_iteration_rows: list[list[object]] = []
    agentic_delta_rows: list[list[object]] = []
    judge_preference_rows: list[list[object]] = []
    aligned_preference_rows: list[list[object]] = []
    confusion_rows: list[list[object]] = []
    detail_rows: list[list[object]] = []
    method_iteration_counts: dict[tuple[str, str, str, str], Counter] = defaultdict(Counter)
    judge_preference_counts: dict[tuple[str, str], Counter] = defaultdict(Counter)
    overall_accuracy_by_model: dict[str, float | None] = {}

    def add_score_stats(
        counter: Counter,
        prediction: int | None,
        allowed: set[int],
        exact: bool,
        abs_diff: int | None,
        disagreed: bool,
    ) -> None:
        counter["cells"] += 1
        counter["exact"] += int(exact)
        counter["invalid"] += int(prediction is None)
        counter["scored"] += int(prediction is not None)
        if abs_diff is not None:
            counter["abs_diff_sum"] += abs_diff
        counter["human_agreed_cells"] += int(not disagreed)
        counter["human_agreed_exact"] += int((not disagreed) and exact)
        counter["human_disagreed_cells"] += int(disagreed)
        counter["human_disagreed_exact"] += int(disagreed and exact)
        if prediction is not None:
            counter[f"llm_score_{prediction}"] += 1
        for human_score in allowed:
            counter[f"human_allowed_{human_score}"] += 1
        if not disagreed:
            counter[f"human_agreed_score_{next(iter(allowed))}"] += 1

    def exact_pct(counter: Counter) -> float | None:
        return pct(counter["exact"], counter["cells"])

    def mean_abs(counter: Counter) -> float | None:
        return pct(counter["abs_diff_sum"], counter["scored"])

    def iteration_sort_value(value: str) -> tuple[int, str]:
        try:
            return (int(float(value)), value)
        except (TypeError, ValueError):
            return (10**9, value)

    for model, predictions_by_key in sorted(predictions_by_model.items()):
        model_counts = Counter()
        all_nearest_pairs: list[tuple[int, int]] = []
        all_item_correct = 0
        bias_counts: dict[tuple[str, str], Counter] = defaultdict(Counter)
        group_counts: dict[tuple[str, str, str, str, str], Counter] = defaultdict(Counter)
        confusion_counts: dict[tuple[str, str, str, str], int] = defaultdict(int)

        for key, predictions in predictions_by_key.items():
            human_item = human_items[key]
            row_correct = []
            first_row = human_item.rows[0] if human_item.rows else {}
            for rubric in rubrics:
                allowed = set(human_item.ratings[rubric])
                if not allowed:
                    continue
                prediction = predictions[rubric]
                disagreed = len(allowed) > 1
                exact = prediction in allowed if prediction is not None else False
                nearest = nearest_human_score(prediction, allowed) if prediction is not None else None
                abs_diff = abs(prediction - nearest) if nearest is not None else None
                agreed_human_score = next(iter(allowed)) if not disagreed else None
                signed_agreed_diff = (
                    prediction - agreed_human_score
                    if prediction is not None and agreed_human_score is not None
                    else None
                )
                model_counts["cells"] += 1
                model_counts["exact"] += int(exact)
                model_counts["invalid_score_cells"] += int(prediction is None)
                model_counts["scored_cells"] += int(prediction is not None)
                if abs_diff is not None:
                    model_counts["abs_diff_sum"] += abs_diff
                model_counts["disagreed_cells"] += int(disagreed)
                model_counts["disagreed_exact"] += int(disagreed and exact)
                model_counts["human_agreed_cells"] += int(not disagreed)
                model_counts["human_agreed_exact"] += int((not disagreed) and exact)
                if nearest is not None and prediction is not None:
                    all_nearest_pairs.append((prediction, nearest))

                bias_key = (model, rubric)
                bias_counts[bias_key]["cells"] += 1
                bias_counts[bias_key]["human_agreed_cells"] += int(not disagreed)
                bias_counts[bias_key]["invalid"] += int(prediction is None)
                bias_counts[bias_key]["exact_on_human_agreed"] += int((not disagreed) and exact)
                if signed_agreed_diff is not None:
                    bias_counts[bias_key]["scored"] += 1
                    bias_counts[bias_key]["signed_agreed_diff_sum"] += signed_agreed_diff
                    bias_counts[bias_key]["above_agreed"] += int(signed_agreed_diff > 0)
                    bias_counts[bias_key]["below_agreed"] += int(signed_agreed_diff < 0)
                    bias_counts[bias_key]["equal_agreed"] += int(signed_agreed_diff == 0)

                group_key = (
                    model,
                    str(first_row.get("model") or ""),
                    str(first_row.get("method") or ""),
                    str(first_row.get("iteration") or ""),
                    rubric,
                )
                group_counts[group_key]["cells"] += 1
                group_counts[group_key]["exact"] += int(exact)
                group_counts[group_key]["human_agreed_cells"] += int(not disagreed)
                group_counts[group_key]["human_agreed_exact"] += int((not disagreed) and exact)
                group_counts[group_key]["human_disagreed_cells"] += int(disagreed)
                group_counts[group_key]["human_disagreed_exact"] += int(disagreed and exact)
                group_counts[group_key]["invalid"] += int(prediction is None)
                if abs_diff is not None:
                    group_counts[group_key]["scored"] += 1
                    group_counts[group_key]["abs_diff_sum"] += abs_diff
                add_score_stats(
                    method_iteration_counts[
                        (
                            model,
                            str(first_row.get("method") or ""),
                            str(first_row.get("iteration") or ""),
                            rubric,
                        )
                    ],
                    prediction,
                    allowed,
                    exact,
                    abs_diff,
                    disagreed,
                )
                add_score_stats(
                    judge_preference_counts[(model, rubric)],
                    prediction,
                    allowed,
                    exact,
                    abs_diff,
                    disagreed,
                )

                human_label = ",".join(str(value) for value in sorted(allowed))
                llm_label = str(prediction) if prediction is not None else "invalid"
                confusion_counts[(model, rubric, human_label, llm_label)] += 1

                row_correct.append(exact)
                detail_rows.append(
                    [
                        model,
                        first_row.get("source_file"),
                        first_row.get("row_number"),
                        first_row.get("chapter_id"),
                        first_row.get("chapter_key"),
                        first_row.get("model"),
                        first_row.get("method"),
                        first_row.get("iteration"),
                        rubric,
                        ",".join(str(value) for value in sorted(allowed)),
                        prediction,
                        nearest,
                        abs_diff,
                        exact,
                        disagreed,
                    ]
                )
            all_item_correct += int(bool(row_correct) and all(row_correct))

        for rubric in rubrics:
            rubric_pairs = []
            counts = Counter()
            for key, predictions in predictions_by_key.items():
                allowed = set(human_items[key].ratings[rubric])
                if not allowed:
                    continue
                prediction = predictions[rubric]
                nearest = nearest_human_score(prediction, allowed) if prediction is not None else None
                abs_diff = abs(prediction - nearest) if nearest is not None else None
                exact = prediction in allowed if prediction is not None else False
                disagreed = len(allowed) > 1
                counts["cells"] += 1
                counts["exact"] += int(exact)
                counts["invalid_score_cells"] += int(prediction is None)
                counts["scored_cells"] += int(prediction is not None)
                if abs_diff is not None:
                    counts["abs_diff_sum"] += abs_diff
                counts["disagreed_cells"] += int(disagreed)
                counts["disagreed_exact"] += int(disagreed and exact)
                counts["human_agreed_cells"] += int(not disagreed)
                counts["human_agreed_exact"] += int((not disagreed) and exact)
                if prediction is not None:
                    counts[f"llm_score_{prediction}"] += 1
                for human_score in allowed:
                    counts[f"human_allowed_{human_score}"] += 1
                if not disagreed:
                    counts[f"human_agreed_score_{next(iter(allowed))}"] += 1
                if prediction is not None and nearest is not None:
                    rubric_pairs.append((prediction, nearest))
            metric_rows.append(
                [
                    model,
                    rubric,
                    counts["cells"],
                    pct(counts["exact"], counts["cells"]),
                    pct(counts["abs_diff_sum"], counts["scored_cells"]),
                    weighted_kappa(rubric_pairs),
                    counts["invalid_score_cells"],
                    counts["human_agreed_cells"],
                    pct(counts["human_agreed_exact"], counts["human_agreed_cells"]),
                    counts["disagreed_cells"],
                    pct(counts["disagreed_exact"], counts["disagreed_cells"]),
                ]
            )
            for score in range(1, 6):
                distribution_rows.append(
                    [
                        model,
                        rubric,
                        score,
                        counts.get(f"llm_score_{score}", 0),
                        pct(counts.get(f"llm_score_{score}", 0), counts["cells"]),
                        counts.get(f"human_allowed_{score}", 0),
                        pct(counts.get(f"human_allowed_{score}", 0), counts["cells"]),
                        counts.get(f"human_agreed_score_{score}", 0),
                        pct(counts.get(f"human_agreed_score_{score}", 0), counts["human_agreed_cells"]),
                    ]
                )

        cells = model_counts["cells"]
        overall_accuracy_by_model[model] = pct(model_counts["exact"], cells)
        summary_rows.append(
            [
                model,
                len(predictions_by_key),
                cells,
                pct(model_counts["exact"], cells),
                pct(model_counts["abs_diff_sum"], model_counts["scored_cells"]),
                weighted_kappa(all_nearest_pairs),
                all_item_correct,
                pct(all_item_correct, len(predictions_by_key)),
                model_counts["invalid_score_cells"],
                model_counts["disagreed_cells"],
                pct(model_counts["disagreed_exact"], model_counts["disagreed_cells"]),
                model_counts["human_agreed_cells"],
                pct(model_counts["human_agreed_exact"], model_counts["human_agreed_cells"]),
                duplicate_stats.get(model, {}).get("duplicate_overlap_items", 0),
                duplicate_stats.get(model, {}).get("conflicting_duplicate_metric_predictions", 0),
            ]
        )

        for (judge_model, rubric), counts in sorted(bias_counts.items()):
            bias_rows.append(
                [
                    judge_model,
                    rubric,
                    counts["cells"],
                    counts["scored"],
                    counts["invalid"],
                    counts["human_agreed_cells"],
                    pct(counts["exact_on_human_agreed"], counts["human_agreed_cells"]),
                    pct(counts["above_agreed"], counts["scored"]),
                    pct(counts["below_agreed"], counts["scored"]),
                    pct(counts["equal_agreed"], counts["scored"]),
                    pct(counts["signed_agreed_diff_sum"], counts["scored"]),
                ]
            )

        for (judge_model, generation_model, method, iteration, rubric), counts in sorted(group_counts.items()):
            group_rows.append(
                [
                    judge_model,
                    generation_model,
                    method,
                    iteration,
                    rubric,
                    counts["cells"],
                    pct(counts["exact"], counts["cells"]),
                    pct(counts["abs_diff_sum"], counts["scored"]),
                    counts["human_agreed_cells"],
                    pct(counts["human_agreed_exact"], counts["human_agreed_cells"]),
                    counts["human_disagreed_cells"],
                    pct(counts["human_disagreed_exact"], counts["human_disagreed_cells"]),
                    counts["invalid"],
                ]
            )

        for (judge_model, rubric, human_scores, llm_label), count in sorted(confusion_counts.items()):
            confusion_rows.append([judge_model, rubric, human_scores, llm_label, count])

    for (judge_model, method, iteration, rubric), counts in sorted(
        method_iteration_counts.items(),
        key=lambda item: (item[0][0], item[0][1], iteration_sort_value(item[0][2]), item[0][3]),
    ):
        row = [
            judge_model,
            method,
            iteration,
            rubric,
            counts["cells"],
            exact_pct(counts),
            mean_abs(counts),
            counts["human_agreed_cells"],
            pct(counts["human_agreed_exact"], counts["human_agreed_cells"]),
            counts["human_disagreed_cells"],
            pct(counts["human_disagreed_exact"], counts["human_disagreed_cells"]),
            counts["invalid"],
        ]
        for score in range(1, 6):
            row.extend(
                [
                    counts.get(f"llm_score_{score}", 0),
                    pct(counts.get(f"llm_score_{score}", 0), counts["cells"]),
                ]
            )
        method_iteration_rows.append(row)

    by_method = defaultdict(dict)
    for key, counts in method_iteration_counts.items():
        judge_model, method, iteration, rubric = key
        by_method[(judge_model, method, rubric)][iteration] = counts
    for (judge_model, method, rubric), by_iteration in sorted(by_method.items()):
        ordered_iterations = sorted(by_iteration, key=iteration_sort_value)
        if len(ordered_iterations) >= 2:
            first_iteration = ordered_iterations[0]
            final_iteration = ordered_iterations[-1]
            first_counts = by_iteration[first_iteration]
            final_counts = by_iteration[final_iteration]
            agentic_delta_rows.append(
                [
                    "final_minus_first",
                    judge_model,
                    method,
                    rubric,
                    first_iteration,
                    final_iteration,
                    first_counts["cells"],
                    final_counts["cells"],
                    exact_pct(first_counts),
                    exact_pct(final_counts),
                    (exact_pct(final_counts) - exact_pct(first_counts))
                    if exact_pct(final_counts) is not None and exact_pct(first_counts) is not None
                    else None,
                    mean_abs(first_counts),
                    mean_abs(final_counts),
                    (mean_abs(final_counts) - mean_abs(first_counts))
                    if mean_abs(final_counts) is not None and mean_abs(first_counts) is not None
                    else None,
                ]
            )

    by_method_delta = defaultdict(dict)
    for key, counts in method_iteration_counts.items():
        judge_model, method, iteration, rubric = key
        by_method_delta[(judge_model, iteration, rubric)][method] = counts
    for (judge_model, iteration, rubric), by_method_name in sorted(
        by_method_delta.items(), key=lambda item: (item[0][0], iteration_sort_value(item[0][1]), item[0][2])
    ):
        if "Holistic" in by_method_name and "Incremental" in by_method_name:
            holistic = by_method_name["Holistic"]
            incremental = by_method_name["Incremental"]
            agentic_delta_rows.append(
                [
                    "incremental_minus_holistic",
                    judge_model,
                    f"iteration_{iteration}",
                    rubric,
                    "Holistic",
                    "Incremental",
                    holistic["cells"],
                    incremental["cells"],
                    exact_pct(holistic),
                    exact_pct(incremental),
                    (exact_pct(incremental) - exact_pct(holistic))
                    if exact_pct(incremental) is not None and exact_pct(holistic) is not None
                    else None,
                    mean_abs(holistic),
                    mean_abs(incremental),
                    (mean_abs(incremental) - mean_abs(holistic))
                    if mean_abs(incremental) is not None and mean_abs(holistic) is not None
                    else None,
                ]
            )

    for (judge_model, rubric), counts in sorted(judge_preference_counts.items()):
        row = [
            judge_model,
            rubric,
            overall_accuracy_by_model.get(judge_model),
            counts["cells"],
            exact_pct(counts),
            mean_abs(counts),
            counts["human_agreed_cells"],
            pct(counts["human_agreed_exact"], counts["human_agreed_cells"]),
            counts["human_disagreed_cells"],
            pct(counts["human_disagreed_exact"], counts["human_disagreed_cells"]),
        ]
        for score in range(1, 6):
            row.extend(
                [
                    counts.get(f"llm_score_{score}", 0),
                    pct(counts.get(f"llm_score_{score}", 0), counts["cells"]),
                    counts.get(f"human_allowed_{score}", 0),
                    pct(counts.get(f"human_allowed_{score}", 0), counts["cells"]),
                    counts.get(f"human_agreed_score_{score}", 0),
                    pct(counts.get(f"human_agreed_score_{score}", 0), counts["human_agreed_cells"]),
                ]
            )
        judge_preference_rows.append(row)

    ranked_models = [
        model
        for model, accuracy in sorted(
            overall_accuracy_by_model.items(),
            key=lambda item: (item[1] is not None, item[1] or -1),
            reverse=True,
        )
    ]
    top_models = set(ranked_models[:3])
    preference_band_counts: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for (judge_model, rubric), counts in judge_preference_counts.items():
        band = "top_3_alignment" if judge_model in top_models else "other_judges"
        dest = preference_band_counts[(band, rubric)]
        for key, value in counts.items():
            dest[key] += value
    for (band, rubric), counts in sorted(preference_band_counts.items()):
        row = [band, rubric, counts["cells"], exact_pct(counts), mean_abs(counts)]
        for score in range(1, 6):
            row.extend(
                [
                    counts.get(f"llm_score_{score}", 0),
                    pct(counts.get(f"llm_score_{score}", 0), counts["cells"]),
                    counts.get(f"human_allowed_{score}", 0),
                    pct(counts.get(f"human_allowed_{score}", 0), counts["cells"]),
                ]
            )
        aligned_preference_rows.append(row)

    summary_rows.sort(key=lambda row: (row[3] is not None, row[3]), reverse=True)
    metric_rows.sort(key=lambda row: (row[0], row[1]))
    distribution_rows.sort(key=lambda row: (row[0], row[1], row[2]))
    bias_rows.sort(key=lambda row: (row[0], row[1]))
    group_rows.sort(key=lambda row: (row[0], row[1], row[2], row[3], row[4]))
    confusion_rows.sort(key=lambda row: (row[0], row[1], row[2], row[3]))
    detail_rows.sort(key=lambda row: (row[0], row[3] or "", row[8]))
    return (
        summary_rows,
        metric_rows,
        distribution_rows,
        bias_rows,
        group_rows,
        method_iteration_rows,
        agentic_delta_rows,
        judge_preference_rows,
        aligned_preference_rows,
        confusion_rows,
        detail_rows,
    )


def write_rows(ws, headers: list[str], rows: Iterable[list[object]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)


def style_workbook(workbook: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(bottom=Side(style="thin", color="D9E2F3"))
    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        for column in range(1, ws.max_column + 1):
            header = display_header(ws.cell(1, column).value)
            width = min(max(len(header) + 2, 12), 44)
            if header in {"human_scores_allowed", "message"}:
                width = 58
            if header in {"chapter_key", "generation_model", "judge_model"}:
                width = 28
            ws.column_dimensions[get_column_letter(column)].width = width
        for row in range(2, ws.max_row + 1):
            for column in range(1, ws.max_column + 1):
                header = normalize_header(ws.cell(1, column).value)
                if "pct" in header or "accuracy" in header:
                    ws.cell(row, column).number_format = "0.0%"
                elif "kappa" in header or "difference" in header or "mean_abs" in header:
                    ws.cell(row, column).number_format = "0.000"


def create_report(
    output_path: Path,
    summary_rows: list[list[object]],
    metric_rows: list[list[object]],
    distribution_rows: list[list[object]],
    bias_rows: list[list[object]],
    group_rows: list[list[object]],
    method_iteration_rows: list[list[object]],
    agentic_delta_rows: list[list[object]],
    judge_preference_rows: list[list[object]],
    aligned_preference_rows: list[list[object]],
    confusion_rows: list[list[object]],
    detail_rows: list[list[object]],
    validation_rows: list[list[object]],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "LLM Human Summary"
    metric = workbook.create_sheet("Metric Summary")
    distributions = workbook.create_sheet("Score Distributions")
    bias = workbook.create_sheet("Score Direction")
    groups = workbook.create_sheet("By Generation Group")
    method_iteration = workbook.create_sheet("Method Iteration Summary")
    agentic_delta = workbook.create_sheet("Agentic Deltas")
    judge_preferences = workbook.create_sheet("Judge Preferences")
    aligned_preferences = workbook.create_sheet("Alignment Preference")
    confusion = workbook.create_sheet("Confusion Matrix")
    details = workbook.create_sheet("Overlap Details")
    validation = workbook.create_sheet("Validation Issues")

    write_rows(
        summary,
        [
            "judge_model",
            "overlap_items",
            "metric_cells",
            "exact_accuracy_any_human",
            "mean_abs_difference_to_allowed",
            "quadratic_weighted_kappa_to_nearest",
            "all_4_metrics_exact_items",
            "all_4_metrics_exact_item_pct",
            "invalid_score_cells",
            "human_disagreed_cells",
            "human_disagreed_exact_accuracy",
            "human_agreed_cells",
            "human_agreed_exact_accuracy",
            "duplicate_overlap_items",
            "conflicting_duplicate_metric_predictions",
        ],
        summary_rows,
    )
    write_rows(
        metric,
        [
            "judge_model",
            "rubric",
            "metric_cells",
            "exact_accuracy_any_human",
            "mean_abs_difference_to_allowed",
            "quadratic_weighted_kappa_to_nearest",
            "invalid_score_cells",
            "human_agreed_cells",
            "human_agreed_exact_accuracy",
            "human_disagreed_cells",
            "human_disagreed_exact_accuracy",
        ],
        metric_rows,
    )
    write_rows(
        distributions,
        [
            "judge_model",
            "rubric",
            "score",
            "llm_score_count",
            "llm_score_pct",
            "human_allowed_count",
            "human_allowed_pct",
            "human_agreed_count",
            "human_agreed_pct",
        ],
        distribution_rows,
    )
    write_rows(
        bias,
        [
            "judge_model",
            "rubric",
            "metric_cells",
            "scored_cells",
            "invalid_score_cells",
            "human_agreed_cells",
            "human_agreed_exact_accuracy",
            "llm_above_agreed_human_pct",
            "llm_below_agreed_human_pct",
            "llm_equal_agreed_human_pct",
            "avg_signed_difference_to_agreed_human",
        ],
        bias_rows,
    )
    write_rows(
        groups,
        [
            "judge_model",
            "generation_model",
            "method",
            "iteration",
            "rubric",
            "metric_cells",
            "exact_accuracy_any_human",
            "mean_abs_difference_to_allowed",
            "human_agreed_cells",
            "human_agreed_exact_accuracy",
            "human_disagreed_cells",
            "human_disagreed_exact_accuracy",
            "invalid_score_cells",
        ],
        group_rows,
    )
    score_pair_headers = []
    for score in range(1, 6):
        score_pair_headers.extend([f"llm_score_{score}_count", f"llm_score_{score}_pct"])
    write_rows(
        method_iteration,
        [
            "judge_model",
            "method",
            "iteration",
            "rubric",
            "metric_cells",
            "exact_accuracy_any_human",
            "mean_abs_difference_to_allowed",
            "human_agreed_cells",
            "human_agreed_exact_accuracy",
            "human_disagreed_cells",
            "human_disagreed_exact_accuracy",
            "invalid_score_cells",
            *score_pair_headers,
        ],
        method_iteration_rows,
    )
    write_rows(
        agentic_delta,
        [
            "comparison",
            "judge_model",
            "method_or_iteration",
            "rubric",
            "baseline_label",
            "comparison_label",
            "baseline_cells",
            "comparison_cells",
            "baseline_exact_accuracy",
            "comparison_exact_accuracy",
            "exact_accuracy_delta",
            "baseline_mean_abs_difference",
            "comparison_mean_abs_difference",
            "mean_abs_difference_delta",
        ],
        agentic_delta_rows,
    )
    preference_headers = [
        "judge_model",
        "rubric",
        "overall_exact_accuracy_any_human",
        "metric_cells",
        "exact_accuracy_any_human",
        "mean_abs_difference_to_allowed",
        "human_agreed_cells",
        "human_agreed_exact_accuracy",
        "human_disagreed_cells",
        "human_disagreed_exact_accuracy",
    ]
    for score in range(1, 6):
        preference_headers.extend(
            [
                f"llm_score_{score}_count",
                f"llm_score_{score}_pct",
                f"human_allowed_{score}_count",
                f"human_allowed_{score}_pct",
                f"human_agreed_{score}_count",
                f"human_agreed_{score}_pct",
            ]
        )
    write_rows(judge_preferences, preference_headers, judge_preference_rows)
    aligned_headers = ["alignment_group", "rubric", "metric_cells", "exact_accuracy_any_human", "mean_abs_difference"]
    for score in range(1, 6):
        aligned_headers.extend(
            [
                f"llm_score_{score}_count",
                f"llm_score_{score}_pct",
                f"human_allowed_{score}_count",
                f"human_allowed_{score}_pct",
            ]
        )
    write_rows(aligned_preferences, aligned_headers, aligned_preference_rows)
    write_rows(
        confusion,
        [
            "judge_model",
            "rubric",
            "human_scores_allowed",
            "llm_score",
            "count",
        ],
        confusion_rows,
    )
    write_rows(
        details,
        [
            "judge_model",
            "human_source_file",
            "human_row_number",
            "chapter_id",
            "chapter_key",
            "generation_model",
            "method",
            "iteration",
            "rubric",
            "human_scores_allowed",
            "llm_score",
            "nearest_human_score",
            "abs_difference_to_allowed",
            "exact_match_any_human",
            "human_disagreed",
        ],
        detail_rows,
    )
    write_rows(
        validation,
        ["source_file", "judge_model", "line_or_row", "issue_type", "message"],
        validation_rows or [["", "", "", "none", "No validation issues found"]],
    )

    workbook["LLM Human Summary"].sheet_properties.tabColor = "1F4E78"
    workbook["Metric Summary"].sheet_properties.tabColor = "5B9BD5"
    workbook["Score Distributions"].sheet_properties.tabColor = "A9D18E"
    workbook["Score Direction"].sheet_properties.tabColor = "9E480E"
    workbook["By Generation Group"].sheet_properties.tabColor = "8064A2"
    workbook["Method Iteration Summary"].sheet_properties.tabColor = "4BACC6"
    workbook["Agentic Deltas"].sheet_properties.tabColor = "F79646"
    workbook["Judge Preferences"].sheet_properties.tabColor = "92D050"
    workbook["Alignment Preference"].sheet_properties.tabColor = "00B050"
    workbook["Confusion Matrix"].sheet_properties.tabColor = "C0504D"
    workbook["Overlap Details"].sheet_properties.tabColor = "70AD47"
    workbook["Validation Issues"].sheet_properties.tabColor = "FFC000"

    style_workbook(workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    base_dir = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(
        description="Compare LLM judge scores with combined human annotations on overlapping rows."
    )
    parser.add_argument(
        "--human-dir",
        type=Path,
        default=base_dir / "results" / "annotation_splits_6",
        help="Directory containing human annotation workbooks.",
    )
    parser.add_argument("--human-glob", default="human_annotation_split_*.xlsx")
    parser.add_argument("--sheet-name", default="Annotation Sample")
    parser.add_argument("--rewritten-column", default="rewritten_text")
    parser.add_argument(
        "--judge-dir",
        type=Path,
        default=base_dir / "results" / "llm_judge",
        help="Directory containing judge_results_*.jsonl files.",
    )
    parser.add_argument("--judge-glob", default="judge_results_*.jsonl")
    parser.add_argument(
        "--output",
        type=Path,
        default=base_dir / "results" / "llm_human_agreement.xlsx",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    human_items, human_issues = read_human_items(
        input_dir=args.human_dir,
        glob_pattern=args.human_glob,
        sheet_name=args.sheet_name,
        rewritten_column=args.rewritten_column,
        rubrics=DEFAULT_RUBRICS,
    )
    predictions_by_model, judge_issues, duplicate_stats = read_llm_predictions(
        judge_dir=args.judge_dir,
        glob_pattern=args.judge_glob,
        human_items=human_items,
        rubrics=DEFAULT_RUBRICS,
    )
    (
        summary_rows,
        metric_rows,
        distribution_rows,
        bias_rows,
        group_rows,
        method_iteration_rows,
        agentic_delta_rows,
        judge_preference_rows,
        aligned_preference_rows,
        confusion_rows,
        detail_rows,
    ) = summarize(
        human_items=human_items,
        predictions_by_model=predictions_by_model,
        duplicate_stats=duplicate_stats,
        rubrics=DEFAULT_RUBRICS,
    )
    create_report(
        output_path=args.output,
        summary_rows=summary_rows,
        metric_rows=metric_rows,
        distribution_rows=distribution_rows,
        bias_rows=bias_rows,
        group_rows=group_rows,
        method_iteration_rows=method_iteration_rows,
        agentic_delta_rows=agentic_delta_rows,
        judge_preference_rows=judge_preference_rows,
        aligned_preference_rows=aligned_preference_rows,
        confusion_rows=confusion_rows,
        detail_rows=detail_rows,
        validation_rows=human_issues + judge_issues,
    )
    print(f"Read {len(human_items)} unique human-overlap items from {args.human_dir}")
    print(f"Compared {len(predictions_by_model)} judge models from {args.judge_dir}")
    print(f"Wrote {args.output}")