# Measure inter-rater agreement on overlapping annotation split workbooks. The split files are treated as rater-specific workbooks. Rows that appear in more than one workbook are overlap items. Ratings must be integers from 1 to 5.
from __future__ import annotations

import argparse
import itertools
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_RUBRICS = [
    "Fidelity",
    "Coverage",
    "Structural Coherence",
    "Readability",
]

DEFAULT_KEY_COLUMNS = [
    "chapter_id",
    "chapter_key",
    "model",
    "method",
    "iteration",
    "original_text",
    "rewritten_text",
]


@dataclass(frozen=True)
class RatingRecord:
    rater: str
    source_file: str
    row_number: int
    row_key: str
    values: dict[str, int | None]
    metadata: dict[str, object]


def normalize_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def display_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def parse_rating(value: object) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        rating = value
    elif isinstance(value, float) and value.is_integer():
        rating = int(value)
    else:
        text = str(value).strip()
        if not re.fullmatch(r"[1-5](?:\.0+)?", text):
            return None
        rating = int(float(text))
    return rating if 1 <= rating <= 5 else None


def safe_text(value: object, max_chars: int = 160) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text[:max_chars]


def find_column(headers: list[object], wanted: str) -> int | None:
    target = normalize_header(wanted)
    for idx, header in enumerate(headers):
        if normalize_header(header) == target:
            return idx
    return None


def row_key_from(row: tuple[object, ...], headers: list[object], key_columns: list[str]) -> str:
    parts = []
    for column in key_columns:
        idx = find_column(headers, column)
        if idx is not None:
            parts.append(f"{column}={safe_text(row[idx], max_chars=500)}")
    return "||".join(parts)


def derive_rater_name(path: Path, pattern: str) -> str:
    match = re.search(pattern, path.stem)
    if match:
        return match.group(1)
    return path.stem


def read_records(
    input_dir: Path,
    glob_pattern: str,
    sheet_name: str,
    rubrics: list[str],
    key_columns: list[str],
    rater_regex: str,
) -> tuple[list[RatingRecord], list[list[object]]]:
    records: list[RatingRecord] = []
    issues: list[list[object]] = []
    files = sorted(input_dir.glob(glob_pattern))
    if not files:
        raise FileNotFoundError(f"No files matched {input_dir / glob_pattern}")

    for path in files:
        workbook = load_workbook(path, data_only=True, read_only=True)
        if sheet_name not in workbook.sheetnames:
            issues.append([path.name, "", "", "missing_sheet", f"Missing sheet: {sheet_name}"])
            continue

        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration:
            issues.append([path.name, "", "", "empty_sheet", "Sheet has no rows"])
            continue

        column_by_rubric = {rubric: find_column(headers, rubric) for rubric in rubrics}
        missing = [rubric for rubric, idx in column_by_rubric.items() if idx is None]
        if missing:
            issues.append([path.name, "", "", "missing_rubric_columns", ", ".join(missing)])
            continue

        key_column_hits = [column for column in key_columns if find_column(headers, column) is not None]
        if not key_column_hits:
            issues.append([path.name, "", "", "missing_key_columns", ", ".join(key_columns)])
            continue

        metadata_columns = {
            "chapter_id": find_column(headers, "chapter_id"),
            "chapter_key": find_column(headers, "chapter_key"),
            "model": find_column(headers, "model"),
            "method": find_column(headers, "method"),
            "iteration": find_column(headers, "iteration"),
        }
        rater = derive_rater_name(path, rater_regex)

        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue

            row_key = row_key_from(row, headers, key_columns)
            values = {}
            for rubric, idx in column_by_rubric.items():
                rating = parse_rating(row[idx]) if idx is not None and idx < len(row) else None
                values[rubric] = rating
                raw_value = row[idx] if idx is not None and idx < len(row) else None
                if raw_value not in (None, "") and rating is None:
                    issues.append([path.name, rater, row_number, "invalid_rating", f"{rubric}: {raw_value!r}"])

            metadata = {
                name: row[idx] if idx is not None and idx < len(row) else None
                for name, idx in metadata_columns.items()
            }
            records.append(
                RatingRecord(
                    rater=rater,
                    source_file=path.name,
                    row_number=row_number,
                    row_key=row_key,
                    values=values,
                    metadata=metadata,
                )
            )

    duplicate_counter = Counter((record.source_file, record.row_key) for record in records)
    for (source_file, row_key), count in duplicate_counter.items():
        if count > 1:
            issues.append([source_file, "", "", "duplicate_row_key", f"{count} rows share key {row_key[:120]}"])

    return records, issues


def pair_records(records: list[RatingRecord]) -> list[tuple[RatingRecord, RatingRecord]]:
    by_key: dict[str, list[RatingRecord]] = defaultdict(list)
    for record in records:
        by_key[record.row_key].append(record)

    pairs = []
    for grouped_records in by_key.values():
        by_rater = {record.rater: record for record in grouped_records}
        if len(by_rater) < 2:
            continue
        for left, right in itertools.combinations(sorted(by_rater.values(), key=lambda item: item.rater), 2):
            pairs.append((left, right))
    return pairs


def weighted_kappa(pairs: list[tuple[int, int]], min_rating: int, max_rating: int) -> float | None:
    if not pairs:
        return None

    categories = list(range(min_rating, max_rating + 1))
    n = len(pairs)
    width = max_rating - min_rating
    left_counts = Counter(left for left, _ in pairs)
    right_counts = Counter(right for _, right in pairs)

    observed = 0.0
    expected = 0.0
    for i in categories:
        for j in categories:
            weight = 0.0 if width == 0 else ((i - j) / width) ** 2
            observed_count = sum(1 for left, right in pairs if left == i and right == j)
            expected_count = left_counts[i] * right_counts[j] / n
            observed += weight * observed_count / n
            expected += weight * expected_count / n

    if math.isclose(expected, 0.0):
        return 1.0 if math.isclose(observed, 0.0) else None
    return 1.0 - observed / expected


def agreement_metrics(pairs: list[tuple[int, int]], min_rating: int, max_rating: int) -> dict[str, object]:
    if not pairs:
        return {
            "paired_scores": 0,
            "exact_agreement_pct": None,
            "mean_abs_difference": None,
            "quadratic_weighted_kappa": None,
        }
    count = len(pairs)
    abs_diffs = [abs(left - right) for left, right in pairs]
    return {
        "paired_scores": count,
        "exact_agreement_pct": sum(diff == 0 for diff in abs_diffs) / count,
        "mean_abs_difference": sum(abs_diffs) / count,
        "quadratic_weighted_kappa": weighted_kappa(pairs, min_rating, max_rating),
    }


def normalized_metadata_value(value: object) -> object:
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def metadata_group_key(record: RatingRecord) -> tuple[object, object, object, object]:
    return (
        normalized_metadata_value(record.metadata.get("chapter_id")),
        record.metadata.get("chapter_key"),
        record.metadata.get("model"),
        record.metadata.get("method"),
    )


def iteration_stage(record: RatingRecord, final_iteration_by_group: dict[tuple[object, object, object, object], int]) -> str:
    iteration = normalized_metadata_value(record.metadata.get("iteration"))
    if iteration == 1:
        return "First"
    if isinstance(iteration, int) and iteration == final_iteration_by_group.get(metadata_group_key(record)):
        return "Final"
    return "Intermediate"


def add_grouped_agreement_rows(
    rows: list[list[object]],
    group_type: str,
    group_value: str,
    grouped_pairs: list[tuple[RatingRecord, RatingRecord]],
    rubrics: list[str],
) -> None:
    all_cell_pairs = []
    for rubric in rubrics:
        score_pairs = [
            (left.values[rubric], right.values[rubric])
            for left, right in grouped_pairs
            if left.values[rubric] is not None and right.values[rubric] is not None
        ]
        all_cell_pairs.extend(score_pairs)
        metrics = agreement_metrics(score_pairs, 1, 5)
        rows.append(
            [
                group_type,
                group_value,
                rubric,
                metrics["paired_scores"],
                metrics["exact_agreement_pct"],
                metrics["mean_abs_difference"],
                metrics["quadratic_weighted_kappa"],
            ]
        )

    metrics = agreement_metrics(all_cell_pairs, 1, 5)
    rows.append(
        [
            group_type,
            group_value,
            "All rubric cells",
            metrics["paired_scores"],
            metrics["exact_agreement_pct"],
            metrics["mean_abs_difference"],
            metrics["quadratic_weighted_kappa"],
        ]
    )


def summarize(
    records: list[RatingRecord],
    pairs: list[tuple[RatingRecord, RatingRecord]],
    rubrics: list[str],
) -> tuple[
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
    list[list[object]],
]:
    summary_rows: list[list[object]] = []
    pair_summary_rows: list[list[object]] = []
    alignment_rows: list[list[object]] = []
    mode_rows: list[list[object]] = []
    detail_rows: list[list[object]] = []
    agreement_by_method_rows: list[list[object]] = []
    agreement_by_iteration_rows: list[list[object]] = []
    agreement_by_stage_rows: list[list[object]] = []
    rating_distribution_rows: list[list[object]] = []
    consensus_preference_rows: list[list[object]] = []

    final_iteration_by_group: dict[tuple[object, object, object, object], int] = {}
    for record in records:
        iteration = normalized_metadata_value(record.metadata.get("iteration"))
        if isinstance(iteration, int):
            key = metadata_group_key(record)
            final_iteration_by_group[key] = max(iteration, final_iteration_by_group.get(key, iteration))

    def add_summary_row(scope: str, rubric: str, score_pairs: list[tuple[int, int]], min_rating: int, max_rating: int):
        metrics = agreement_metrics(score_pairs, min_rating, max_rating)
        summary_rows.append(
            [
                scope,
                rubric,
                metrics["paired_scores"],
                metrics["exact_agreement_pct"],
                metrics["mean_abs_difference"],
                metrics["quadratic_weighted_kappa"],
            ]
        )

    all_cell_pairs = []
    for rubric in rubrics:
        score_pairs = [
            (left.values[rubric], right.values[rubric])
            for left, right in pairs
            if left.values[rubric] is not None and right.values[rubric] is not None
        ]
        all_cell_pairs.extend(score_pairs)
        add_summary_row("All raters", rubric, score_pairs, 1, 5)

    add_summary_row("All raters", "All rubric cells", all_cell_pairs, 1, 5)

    all_metric_match_count = 0
    exact_cell_count_distribution = Counter()
    for left, right in pairs:
        exact_count = 0
        comparable_count = 0
        abs_diffs = []
        for rubric in rubrics:
            left_value = left.values[rubric]
            right_value = right.values[rubric]
            if left_value is None or right_value is None:
                continue
            comparable_count += 1
            diff = abs(left_value - right_value)
            abs_diffs.append(diff)
            exact_count += int(diff == 0)
        all_exact = comparable_count == len(rubrics) and exact_count == len(rubrics)
        all_metric_match_count += int(all_exact)
        exact_cell_count_distribution[exact_count] += 1
        alignment_rows.append(
            [
                left.rater,
                right.rater,
                left.metadata.get("chapter_id"),
                left.metadata.get("chapter_key"),
                left.metadata.get("model"),
                left.metadata.get("method"),
                left.metadata.get("iteration"),
                iteration_stage(left, final_iteration_by_group),
                comparable_count,
                exact_count,
                exact_count / comparable_count if comparable_count else None,
                all_exact,
                sum(abs_diffs) / len(abs_diffs) if abs_diffs else None,
                max(abs_diffs) if abs_diffs else None,
            ]
        )

    summary_rows.append(
        [
            "All raters",
            "All rubrics aligned per item",
            len(pairs),
            all_metric_match_count / len(pairs) if pairs else None,
            None,
            None,
        ]
    )
    for exact_count in range(len(rubrics) + 1):
        summary_rows.append(
            [
                "All raters",
                f"Items with {exact_count} exact rubric cells",
                exact_cell_count_distribution.get(exact_count, 0),
                exact_cell_count_distribution.get(exact_count, 0) / len(pairs) if pairs else None,
                None,
                None,
            ]
        )

    by_key: dict[str, list[RatingRecord]] = defaultdict(list)
    for record in records:
        by_key[record.row_key].append(record)

    for row_key, grouped_records in sorted(by_key.items()):
        if len({record.rater for record in grouped_records}) < 2:
            continue
        first = grouped_records[0]
        for rubric in rubrics:
            values = [record.values[rubric] for record in grouped_records if record.values[rubric] is not None]
            if not values:
                continue
            counts = Counter(values)
            max_count = max(counts.values())
            modes = sorted(value for value, count in counts.items() if count == max_count)
            mode_rows.append(
                [
                    first.metadata.get("chapter_id"),
                    first.metadata.get("chapter_key"),
                    first.metadata.get("model"),
                    first.metadata.get("method"),
                    first.metadata.get("iteration"),
                    iteration_stage(first, final_iteration_by_group),
                    rubric,
                    len(values),
                    ",".join(str(value) for value in values),
                    ",".join(str(value) for value in modes),
                    max_count,
                    len(set(values)) > 1,
                    min(values),
                    max(values),
                    max(values) - min(values),
                    sum(values) / len(values),
                ]
            )

    by_rater_pair: dict[tuple[str, str], list[tuple[RatingRecord, RatingRecord]]] = defaultdict(list)
    for left, right in pairs:
        by_rater_pair[(left.rater, right.rater)].append((left, right))

    for (left_rater, right_rater), rater_pairs in sorted(by_rater_pair.items()):
        for rubric in rubrics:
            score_pairs = [
                (left.values[rubric], right.values[rubric])
                for left, right in rater_pairs
                if left.values[rubric] is not None and right.values[rubric] is not None
            ]
            metrics = agreement_metrics(score_pairs, 1, 5)
            pair_summary_rows.append(
                [
                    left_rater,
                    right_rater,
                    rubric,
                    metrics["paired_scores"],
                    metrics["exact_agreement_pct"],
                    metrics["mean_abs_difference"],
                    metrics["quadratic_weighted_kappa"],
                ]
            )

    for left, right in sorted(pairs, key=lambda item: (item[0].rater, item[1].rater, item[0].row_key)):
        left_ratings = [left.values[rubric] for rubric in rubrics]
        right_ratings = [right.values[rubric] for rubric in rubrics]
        row = [
            left.rater,
            right.rater,
            left.metadata.get("chapter_id"),
            left.metadata.get("chapter_key"),
            left.metadata.get("model"),
            left.metadata.get("method"),
            left.metadata.get("iteration"),
            iteration_stage(left, final_iteration_by_group),
        ]
        for rubric in rubrics:
            left_value = left.values[rubric]
            right_value = right.values[rubric]
            row.extend(
                [
                    left_value,
                    right_value,
                    abs(left_value - right_value)
                    if left_value is not None and right_value is not None
                    else None,
                ]
            )
        detail_rows.append(row)

    by_method: dict[str, list[tuple[RatingRecord, RatingRecord]]] = defaultdict(list)
    by_iteration: dict[str, list[tuple[RatingRecord, RatingRecord]]] = defaultdict(list)
    by_stage: dict[str, list[tuple[RatingRecord, RatingRecord]]] = defaultdict(list)
    by_method_stage: dict[str, list[tuple[RatingRecord, RatingRecord]]] = defaultdict(list)
    for left, right in pairs:
        method = str(left.metadata.get("method") or "")
        iteration = str(normalized_metadata_value(left.metadata.get("iteration")) or "")
        stage = iteration_stage(left, final_iteration_by_group)
        by_method[method].append((left, right))
        by_iteration[iteration].append((left, right))
        by_stage[stage].append((left, right))
        by_method_stage[f"{method} | {stage}"].append((left, right))

    for method, grouped_pairs in sorted(by_method.items()):
        add_grouped_agreement_rows(agreement_by_method_rows, "method", method, grouped_pairs, rubrics)
    for iteration, grouped_pairs in sorted(by_iteration.items(), key=lambda item: int(item[0]) if item[0].isdigit() else 999):
        add_grouped_agreement_rows(agreement_by_iteration_rows, "iteration", iteration, grouped_pairs, rubrics)
    for stage in ["First", "Intermediate", "Final"]:
        if stage in by_stage:
            add_grouped_agreement_rows(agreement_by_stage_rows, "iteration_stage", stage, by_stage[stage], rubrics)
    for group_value, grouped_pairs in sorted(by_method_stage.items()):
        add_grouped_agreement_rows(agreement_by_stage_rows, "method_iteration_stage", group_value, grouped_pairs, rubrics)

    overlap_keys = {left.row_key for left, _right in pairs} | {right.row_key for _left, right in pairs}
    overlap_records = [record for record in records if record.row_key in overlap_keys]

    def add_distribution_row(scope_type: str, scope_value: str, rubric: str, values: list[int]) -> None:
        counts = Counter(values)
        total = sum(counts.values())
        row = [scope_type, scope_value, rubric, total]
        row.extend(counts.get(score, 0) for score in range(1, 6))
        row.extend((counts.get(score, 0) / total if total else None) for score in range(1, 6))
        row.extend(
            [
                (counts.get(1, 0) + counts.get(2, 0)) / total if total else None,
                counts.get(3, 0) / total if total else None,
                (counts.get(4, 0) + counts.get(5, 0)) / total if total else None,
            ]
        )
        rating_distribution_rows.append(row)

    distribution_groups: dict[tuple[str, str], list[RatingRecord]] = defaultdict(list)
    distribution_groups[("all_overlap", "All overlap ratings")] = overlap_records
    for record in overlap_records:
        method = str(record.metadata.get("method") or "")
        iteration = str(normalized_metadata_value(record.metadata.get("iteration")) or "")
        stage = iteration_stage(record, final_iteration_by_group)
        distribution_groups[("method", method)].append(record)
        distribution_groups[("iteration", iteration)].append(record)
        distribution_groups[("iteration_stage", stage)].append(record)
        distribution_groups[("method_iteration_stage", f"{method} | {stage}")].append(record)

    for (scope_type, scope_value), grouped_records in sorted(distribution_groups.items()):
        for rubric in rubrics:
            values = [record.values[rubric] for record in grouped_records if record.values[rubric] is not None]
            add_distribution_row(scope_type, scope_value, rubric, values)

    mode_distribution_groups: dict[tuple[str, str], list[list[object]]] = defaultdict(list)
    for row in mode_rows:
        method = str(row[3] or "")
        iteration = str(normalized_metadata_value(row[4]) or "")
        stage = str(row[5] or "")
        mode_distribution_groups[("all_overlap", "All consensus rows")].append(row)
        mode_distribution_groups[("method", method)].append(row)
        mode_distribution_groups[("iteration", iteration)].append(row)
        mode_distribution_groups[("iteration_stage", stage)].append(row)
        mode_distribution_groups[("method_iteration_stage", f"{method} | {stage}")].append(row)

    for (scope_type, scope_value), grouped_rows in sorted(mode_distribution_groups.items()):
        for rubric in rubrics:
            rubric_rows = [row for row in grouped_rows if row[6] == rubric]
            value_weights = Counter()
            tie_count = 0
            disagreed_count = 0
            for row in rubric_rows:
                modes = [int(value) for value in str(row[9]).split(",") if value]
                if len(modes) > 1:
                    tie_count += 1
                if row[11]:
                    disagreed_count += 1
                for value in modes:
                    value_weights[value] += 1 / len(modes)
            total = len(rubric_rows)
            preference_total = sum(value_weights.values())
            preference_values = [value_weights.get(score, 0) for score in range(1, 6)]
            consensus_preference_rows.append(
                [
                    scope_type,
                    scope_value,
                    rubric,
                    total,
                    disagreed_count,
                    disagreed_count / total if total else None,
                    tie_count,
                    tie_count / total if total else None,
                    *preference_values,
                    *[
                        value / preference_total if preference_total else None
                        for value in preference_values
                    ],
                ]
            )

    return (
        summary_rows,
        pair_summary_rows,
        alignment_rows,
        mode_rows,
        detail_rows,
        agreement_by_method_rows,
        agreement_by_iteration_rows,
        agreement_by_stage_rows,
        rating_distribution_rows,
        consensus_preference_rows,
    )


def write_rows(ws, headers: list[str], rows: Iterable[list[object]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)


def style_workbook(workbook: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(bottom=Side(style="thin", color="D9E2F3"))

    for ws in workbook.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=False)

        for column in range(1, ws.max_column + 1):
            letter = get_column_letter(column)
            header = display_header(ws.cell(1, column).value)
            width = min(max(len(header) + 2, 12), 42)
            if header in {"row_key", "message"}:
                width = 60
            if header in {"chapter_key", "model"}:
                width = 26
            ws.column_dimensions[letter].width = width

        for row in range(2, ws.max_row + 1):
            for column in range(1, ws.max_column + 1):
                header = normalize_header(ws.cell(1, column).value)
                if "pct" in header:
                    ws.cell(row, column).number_format = "0.0%"
                elif "kappa" in header or "difference" in header:
                    ws.cell(row, column).number_format = "0.000"


def create_report(
    output_path: Path,
    summary_rows: list[list[object]],
    pair_summary_rows: list[list[object]],
    alignment_rows: list[list[object]],
    mode_rows: list[list[object]],
    detail_rows: list[list[object]],
    agreement_by_method_rows: list[list[object]],
    agreement_by_iteration_rows: list[list[object]],
    agreement_by_stage_rows: list[list[object]],
    rating_distribution_rows: list[list[object]],
    consensus_preference_rows: list[list[object]],
    validation_rows: list[list[object]],
    rubrics: list[str],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Agreement Summary"
    rater_pairs = workbook.create_sheet("Rater Pair Summary")
    by_method = workbook.create_sheet("Agreement by Method")
    by_iteration = workbook.create_sheet("Agreement by Iteration")
    by_stage = workbook.create_sheet("Agreement by Stage")
    distributions = workbook.create_sheet("Rating Distributions")
    consensus = workbook.create_sheet("Consensus Preferences")
    alignment = workbook.create_sheet("Overall Alignment")
    modes = workbook.create_sheet("Rubric Modes")
    details = workbook.create_sheet("Overlap Details")
    validation = workbook.create_sheet("Validation Issues")

    write_rows(
        summary,
        [
            "scope",
            "rubric",
            "paired_scores",
            "exact_agreement_pct",
            "mean_abs_difference",
            "quadratic_weighted_kappa",
        ],
        summary_rows,
    )
    write_rows(
        rater_pairs,
        [
            "rater_left",
            "rater_right",
            "rubric",
            "paired_scores",
            "exact_agreement_pct",
            "mean_abs_difference",
            "quadratic_weighted_kappa",
        ],
        pair_summary_rows,
    )
    grouped_agreement_headers = [
        "group_type",
        "group_value",
        "rubric",
        "paired_scores",
        "exact_agreement_pct",
        "mean_abs_difference",
        "quadratic_weighted_kappa",
    ]
    write_rows(by_method, grouped_agreement_headers, agreement_by_method_rows)
    write_rows(by_iteration, grouped_agreement_headers, agreement_by_iteration_rows)
    write_rows(by_stage, grouped_agreement_headers, agreement_by_stage_rows)
    distribution_headers = [
        "scope_type",
        "scope_value",
        "rubric",
        "rating_count",
        "rating_1_count",
        "rating_2_count",
        "rating_3_count",
        "rating_4_count",
        "rating_5_count",
        "rating_1_pct",
        "rating_2_pct",
        "rating_3_pct",
        "rating_4_pct",
        "rating_5_pct",
        "ratings_1_2_pct",
        "rating_3_pct_mid",
        "ratings_4_5_pct",
    ]
    write_rows(distributions, distribution_headers, rating_distribution_rows)
    consensus_headers = [
        "scope_type",
        "scope_value",
        "rubric",
        "consensus_rows",
        "disagreed_rows",
        "disagreed_pct",
        "tied_mode_rows",
        "tied_mode_pct",
        "mode_1_weighted_count",
        "mode_2_weighted_count",
        "mode_3_weighted_count",
        "mode_4_weighted_count",
        "mode_5_weighted_count",
        "mode_1_pct",
        "mode_2_pct",
        "mode_3_pct",
        "mode_4_pct",
        "mode_5_pct",
    ]
    write_rows(consensus, consensus_headers, consensus_preference_rows)
    write_rows(
        alignment,
        [
            "rater_left",
            "rater_right",
            "chapter_id",
            "chapter_key",
            "model",
            "method",
            "iteration",
            "iteration_stage",
            "comparable_rubric_cells",
            "exact_rubric_cells",
            "exact_rubric_cell_pct",
            "all_rubrics_exact",
            "mean_abs_difference",
            "max_abs_difference",
        ],
        alignment_rows,
    )
    write_rows(
        modes,
        [
            "chapter_id",
            "chapter_key",
            "model",
            "method",
            "iteration",
            "iteration_stage",
            "rubric",
            "human_rating_count",
            "human_scores",
            "mode_score",
            "mode_count",
            "human_disagreed",
            "min_score",
            "max_score",
            "score_range",
            "mean_score",
        ],
        mode_rows,
    )

    detail_headers = [
        "rater_left",
        "rater_right",
        "chapter_id",
        "chapter_key",
        "model",
        "method",
        "iteration",
        "iteration_stage",
    ]
    for rubric in rubrics:
        safe_rubric = rubric.lower().replace(" ", "_")
        detail_headers.extend(
            [
                f"{safe_rubric}_left",
                f"{safe_rubric}_right",
                f"{safe_rubric}_abs_diff",
            ]
        )
    write_rows(details, detail_headers, detail_rows)
    write_rows(
        validation,
        ["source_file", "rater", "row_number", "issue_type", "message"],
        validation_rows or [["", "", "", "none", "No validation issues found"]],
    )

    workbook["Agreement Summary"].sheet_properties.tabColor = "1F4E78"
    workbook["Rater Pair Summary"].sheet_properties.tabColor = "5B9BD5"
    workbook["Agreement by Method"].sheet_properties.tabColor = "4472C4"
    workbook["Agreement by Iteration"].sheet_properties.tabColor = "4472C4"
    workbook["Agreement by Stage"].sheet_properties.tabColor = "4472C4"
    workbook["Rating Distributions"].sheet_properties.tabColor = "A9D18E"
    workbook["Consensus Preferences"].sheet_properties.tabColor = "A9D18E"
    workbook["Overall Alignment"].sheet_properties.tabColor = "70AD47"
    workbook["Rubric Modes"].sheet_properties.tabColor = "A9D18E"
    workbook["Overlap Details"].sheet_properties.tabColor = "70AD47"
    workbook["Validation Issues"].sheet_properties.tabColor = "FFC000"

    style_workbook(workbook)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Calculate inter-rater agreement for overlapping annotation split files."
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "results" / "annotation_splits_6",
        help="Directory containing completed split workbooks.",
    )
    parser.add_argument(
        "--glob",
        default="human_annotation_split_*.xlsx",
        help="Workbook filename pattern inside --input-dir.",
    )
    parser.add_argument(
        "--sheet-name",
        default="Annotation Sample",
        help="Sheet containing the annotation rows.",
    )
    parser.add_argument(
        "--rubrics",
        nargs="+",
        default=DEFAULT_RUBRICS,
        help="Rubric columns scored from 1 to 5. Matching ignores extra whitespace and case.",
    )
    parser.add_argument(
        "--key-columns",
        nargs="+",
        default=DEFAULT_KEY_COLUMNS,
        help="Columns used to identify the same annotation row across split files.",
    )
    parser.add_argument(
        "--rater-regex",
        default=r"split_(\d+)",
        help="Regex with one capture group used to derive rater names from filenames.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "results" / "inter_rater_agreement.xlsx",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    records, validation_rows = read_records(
        input_dir=args.input_dir,
        glob_pattern=args.glob,
        sheet_name=args.sheet_name,
        rubrics=args.rubrics,
        key_columns=args.key_columns,
        rater_regex=args.rater_regex,
    )
    pairs = pair_records(records)
    (
        summary_rows,
        pair_summary_rows,
        alignment_rows,
        mode_rows,
        detail_rows,
        agreement_by_method_rows,
        agreement_by_iteration_rows,
        agreement_by_stage_rows,
        rating_distribution_rows,
        consensus_preference_rows,
    ) = summarize(records, pairs, args.rubrics)
    create_report(
        args.output,
        summary_rows,
        pair_summary_rows,
        alignment_rows,
        mode_rows,
        detail_rows,
        agreement_by_method_rows,
        agreement_by_iteration_rows,
        agreement_by_stage_rows,
        rating_distribution_rows,
        consensus_preference_rows,
        validation_rows,
        args.rubrics,
    )
    print(f"Read {len(records)} annotated rows from {args.input_dir}")
    print(f"Found {len(pairs)} overlap item pairs")
    print(f"Wrote {args.output}")