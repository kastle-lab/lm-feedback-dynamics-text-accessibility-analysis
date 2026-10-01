# Generate a deeper failure-analysis workbook from DysText result logs

from __future__ import annotations

import argparse
import ast
import csv
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


RESULT_SETS = [
    {"mode": "Holistic", "root": "results/local_holistic", "pattern": "*_agent_log.csv"},
    {"mode": "Incremental", "root": "results/local_incremental", "pattern": "*_agent_log.csv"},
    {"mode": "Holistic", "model": "gpt-global", "path": "results/new_agent_log_hollistic.csv"},
    {"mode": "Incremental", "model": "gpt-global", "path": "results/all_agent_log_incremental.csv"},
]


OVERVIEW_HEADERS = [
    "Model",
    "Mode",
    "Rows",
    "Chapters",
    "Passed Iterations",
    "Pass Rate",
    "Factual Drift Failures",
    "Factual Drift Rate",
    "Longphrase Failures",
    "Most Common Failed Check",
    "Most Common Failed Check Count",
    "Worst Drift Iteration",
    "Worst Drift Iteration Count",
    "Best Drift Iteration",
    "Best Drift Iteration Count",
]

ITERATION_HEADERS = [
    "Model",
    "Mode",
    "Iteration",
    "Rows",
    "Passes",
    "Pass Rate",
    "Factual Drift Failures",
    "Factual Drift Rate",
    "Longphrase Failures",
    "Total Failed Checks",
]

CRITERION_HEADERS = [
    "Model",
    "Mode",
    "Failed Check",
    "Count",
    "Share of Rows",
    "Share of Failed Checks",
]

DELTA_HEADERS = [
    "Model",
    "Holistic Drift Failures",
    "Incremental Drift Failures",
    "Incremental Minus Holistic",
    "Holistic Drift Rate",
    "Incremental Drift Rate",
    "Holistic Pass Rate",
    "Incremental Pass Rate",
    "Holistic Worst Iteration",
    "Incremental Worst Iteration",
    "Holistic Worst Iteration Count",
    "Incremental Worst Iteration Count",
    "Holistic Best Iteration",
    "Incremental Best Iteration",
    "Holistic Best Iteration Count",
    "Incremental Best Iteration Count",
]

HOTSPOT_HEADERS = [
    "Model",
    "Mode",
    "Chapter Key",
    "Factual Drift Failures",
    "Failed Iterations",
    "Total Failed Checks",
    "Passes",
    "Rows",
    "Source Log",
]

DEFINITION_ROWS = [
    ["Factual Drift Rate", "factual_drift failures divided by logged rows for that model and mode."],
    ["Worst Drift Iteration", "The iteration number with the highest factual_drift count."],
    ["Best Drift Iteration", "The iteration number with the lowest factual_drift count."],
    ["Criterion Breakdown", "Counts every label in failed_checks. One row can contribute multiple labels."],
    ["Method Delta", "Incremental minus holistic factual drift. Positive means incremental drifted more."],
    ["Chapter Hotspots", "Chapters with the most factual_drift failures within a model and mode."],
]


def as_bool(value: object) -> bool:
    return str(value).strip().lower() == "true"


def labels(value: object) -> list[str]:
    if not value:
        return []
    try:
        parsed = ast.literal_eval(str(value))
        if isinstance(parsed, list):
            return [str(item) for item in parsed]
    except (SyntaxError, ValueError):
        pass
    return [item.strip() for item in str(value).split(",") if item.strip()]


def model_name_from_log(path: Path) -> str:
    name = path.name
    if name.endswith("_agent_log.csv"):
        return name[: -len("_agent_log.csv")]
    return name


def read_rows(path: Path) -> list[dict[str, str]]:
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def iter_logs(base_dir: Path):
    for spec in RESULT_SETS:
        mode = spec["mode"]
        if "path" in spec:
            path = base_dir / spec["path"]
            yield mode, spec.get("model") or model_name_from_log(path), path
            continue
        root = base_dir / spec["root"]
        for path in sorted(root.glob(spec["pattern"])):
            yield mode, model_name_from_log(path), path


def annotate_rows(base_dir: Path) -> list[dict[str, object]]:
    annotated = []
    for mode, model, path in iter_logs(base_dir):
        rows = read_rows(path)
        for index, row in enumerate(rows):
            chapter_index = index // 4 + 1
            row_labels = labels(row.get("failed_checks"))
            annotated.append(
                {
                    "Model": model,
                    "Mode": mode,
                    "Source Log": str(path),
                    "Chapter Index": chapter_index,
                    "Chapter Key": row.get("chapter_key", ""),
                    "Iteration": int(row.get("iteration") or (index % 4 + 1)),
                    "Reached Judge": as_bool(row.get("reached_judge")),
                    "Labels": row_labels,
                    "Factual Drift": "factual_drift" in row_labels,
                    "Longphrase": "longphrase" in row_labels,
                }
            )
    return annotated


def pct(numerator: int | float, denominator: int | float) -> float:
    return numerator / denominator if denominator else 0


def most_common(counter: Counter) -> tuple[str, int]:
    if not counter:
        return "", 0
    return counter.most_common(1)[0]


def drift_iteration_extremes(group: list[dict[str, object]]) -> tuple[int | str, int, int | str, int]:
    iterations = sorted({int(row["Iteration"]) for row in group})
    if not iterations:
        return "", 0, "", 0

    drift_by_iteration = Counter(row["Iteration"] for row in group if row["Factual Drift"])
    counts = [(iteration, drift_by_iteration.get(iteration, 0)) for iteration in iterations]
    worst_iteration, worst_count = max(counts, key=lambda item: (item[1], -item[0]))
    best_iteration, best_count = min(counts, key=lambda item: (item[1], item[0]))
    return worst_iteration, worst_count, best_iteration, best_count


def overview_rows(rows: list[dict[str, object]]) -> list[list[object]]:
    groups: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        groups[(str(row["Model"]), str(row["Mode"]))].append(row)

    output = []
    for (model, mode), group in sorted(groups.items()):
        label_counts = Counter(label for row in group for label in row["Labels"])
        common_label, common_count = most_common(label_counts)
        worst_iteration, worst_count, best_iteration, best_count = drift_iteration_extremes(group)
        drift_count = sum(1 for row in group if row["Factual Drift"])
        pass_count = sum(1 for row in group if row["Reached Judge"])
        longphrase_count = sum(1 for row in group if row["Longphrase"])
        output.append(
            [
                model,
                mode,
                len(group),
                len({row["Chapter Index"] for row in group}),
                pass_count,
                pct(pass_count, len(group)),
                drift_count,
                pct(drift_count, len(group)),
                longphrase_count,
                common_label,
                common_count,
                worst_iteration,
                worst_count,
                best_iteration,
                best_count,
            ]
        )
    return output


def iteration_rows(rows: list[dict[str, object]]) -> list[list[object]]:
    groups: dict[tuple[str, str, int], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        groups[(str(row["Model"]), str(row["Mode"]), int(row["Iteration"]))].append(row)

    output = []
    for (model, mode, iteration), group in sorted(groups.items()):
        passes = sum(1 for row in group if row["Reached Judge"])
        drift = sum(1 for row in group if row["Factual Drift"])
        longphrase = sum(1 for row in group if row["Longphrase"])
        failed_checks = sum(len(row["Labels"]) for row in group)
        output.append(
            [
                model,
                mode,
                iteration,
                len(group),
                passes,
                pct(passes, len(group)),
                drift,
                pct(drift, len(group)),
                longphrase,
                failed_checks,
            ]
        )
    return output


def criterion_rows(rows: list[dict[str, object]]) -> list[list[object]]:
    groups: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        groups[(str(row["Model"]), str(row["Mode"]))].append(row)

    output = []
    for (model, mode), group in sorted(groups.items()):
        label_counts = Counter(label for row in group for label in row["Labels"])
        total_labels = sum(label_counts.values())
        for label, count in sorted(label_counts.items(), key=lambda item: (-item[1], item[0])):
            output.append(
                [
                    model,
                    mode,
                    label,
                    count,
                    pct(count, len(group)),
                    pct(count, total_labels),
                ]
            )
    return output


def method_delta_rows(rows: list[dict[str, object]]) -> list[list[object]]:
    by_model_mode: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        by_model_mode[(str(row["Model"]), str(row["Mode"]))].append(row)

    models = sorted({str(row["Model"]) for row in rows})
    output = []
    for model in models:
        mode_data = {}
        for mode in ["Holistic", "Incremental"]:
            group = by_model_mode.get((model, mode), [])
            worst_iteration, worst_count, best_iteration, best_count = drift_iteration_extremes(group)
            drift_count = sum(1 for row in group if row["Factual Drift"])
            pass_count = sum(1 for row in group if row["Reached Judge"])
            mode_data[mode] = {
                "rows": len(group),
                "drift": drift_count,
                "drift_rate": pct(drift_count, len(group)),
                "pass_rate": pct(pass_count, len(group)),
                "worst_iteration": worst_iteration,
                "worst_count": worst_count,
                "best_iteration": best_iteration,
                "best_count": best_count,
            }
        output.append(
            [
                model,
                mode_data["Holistic"]["drift"],
                mode_data["Incremental"]["drift"],
                mode_data["Incremental"]["drift"] - mode_data["Holistic"]["drift"],
                mode_data["Holistic"]["drift_rate"],
                mode_data["Incremental"]["drift_rate"],
                mode_data["Holistic"]["pass_rate"],
                mode_data["Incremental"]["pass_rate"],
                mode_data["Holistic"]["worst_iteration"],
                mode_data["Incremental"]["worst_iteration"],
                mode_data["Holistic"]["worst_count"],
                mode_data["Incremental"]["worst_count"],
                mode_data["Holistic"]["best_iteration"],
                mode_data["Incremental"]["best_iteration"],
                mode_data["Holistic"]["best_count"],
                mode_data["Incremental"]["best_count"],
            ]
        )
    return output


def hotspot_rows(rows: list[dict[str, object]], top_n: int) -> list[list[object]]:
    groups: dict[tuple[str, str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        groups[(str(row["Model"]), str(row["Mode"]), str(row["Chapter Key"]))].append(row)

    by_model_mode: dict[tuple[str, str], list[list[object]]] = defaultdict(list)
    for (model, mode, chapter_key), group in groups.items():
        drift_iterations = [
            str(row["Iteration"]) for row in group if row["Factual Drift"]
        ]
        if not drift_iterations:
            continue
        by_model_mode[(model, mode)].append(
            [
                model,
                mode,
                chapter_key,
                len(drift_iterations),
                ", ".join(drift_iterations),
                sum(len(row["Labels"]) for row in group),
                sum(1 for row in group if row["Reached Judge"]),
                len(group),
                group[0]["Source Log"],
            ]
        )

    output = []
    for key in sorted(by_model_mode):
        output.extend(
            sorted(
                by_model_mode[key],
                key=lambda row: (-int(row[3]), str(row[2])),
            )[:top_n]
        )
    return output


def write_title(ws, title: str, subtitle: str, cols: int):
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True, color="1F1F1F")
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="666666")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(cols, 6))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(cols, 8))


def add_table(ws, headers: list[str], rows: list[list[object]]):
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=3, column=col_idx, value=header)
    for row_idx, row in enumerate(rows, start=4):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def style_sheet(ws, row_count: int, col_count: int, percent_cols: set[int] | None = None):
    percent_cols = percent_cols or set()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin_gray)

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    for cell in ws[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=4, max_row=max(row_count, 4), max_col=col_count):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if cell.column in percent_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%"
            elif isinstance(cell.value, float):
                cell.number_format = "0.00"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"

    for col_idx in range(1, col_count + 1):
        letter = get_column_letter(col_idx)
        max_len = max(
            len(str(ws.cell(row=row, column=col_idx).value or ""))
            for row in range(1, row_count + 1)
        )
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)


def add_iteration_chart(ws, row_count: int):
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Factual drift by iteration"
    chart.y_axis.title = "Failures"
    chart.x_axis.title = "Model and mode rows"
    data = Reference(ws, min_col=7, min_row=3, max_row=row_count)
    cats = Reference(ws, min_col=3, min_row=4, max_row=row_count)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 18
    ws.add_chart(chart, "L3")


def add_delta_chart(ws, row_count: int):
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Incremental minus holistic drift"
    chart.y_axis.title = "Model"
    chart.x_axis.title = "Difference"
    data = Reference(ws, min_col=4, min_row=3, max_row=row_count)
    cats = Reference(ws, min_col=1, min_row=4, max_row=row_count)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 16
    ws.add_chart(chart, "L3")


def write_workbook(rows: list[dict[str, object]], output_path: Path, top_n: int):
    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"
    iteration = wb.create_sheet("By Iteration")
    criterion = wb.create_sheet("By Failed Check")
    delta = wb.create_sheet("Method Delta")
    hotspots = wb.create_sheet("Chapter Hotspots")
    definitions = wb.create_sheet("Definitions")

    overview_data = overview_rows(rows)
    write_title(overview, "Failure analysis overview", "One row per model and feedback method.", len(OVERVIEW_HEADERS))
    add_table(overview, OVERVIEW_HEADERS, overview_data)
    style_sheet(overview, len(overview_data) + 3, len(OVERVIEW_HEADERS), {6, 8})

    iteration_data = iteration_rows(rows)
    write_title(iteration, "Failures by iteration", "Shows where factual drift and other failures concentrate over the four passes.", len(ITERATION_HEADERS))
    add_table(iteration, ITERATION_HEADERS, iteration_data)
    style_sheet(iteration, len(iteration_data) + 3, len(ITERATION_HEADERS), {6, 8})
    add_iteration_chart(iteration, len(iteration_data) + 3)

    criterion_data = criterion_rows(rows)
    write_title(criterion, "Failed-check breakdown", "Counts each failed-check label by model and method.", len(CRITERION_HEADERS))
    add_table(criterion, CRITERION_HEADERS, criterion_data)
    style_sheet(criterion, len(criterion_data) + 3, len(CRITERION_HEADERS), {5, 6})

    delta_data = method_delta_rows(rows)
    write_title(delta, "Holistic vs incremental", "Positive delta means incremental produced more factual drift failures.", len(DELTA_HEADERS))
    add_table(delta, DELTA_HEADERS, delta_data)
    style_sheet(delta, len(delta_data) + 3, len(DELTA_HEADERS), {5, 6, 7, 8})
    add_delta_chart(delta, len(delta_data) + 3)

    hotspot_data = hotspot_rows(rows, top_n)
    write_title(hotspots, "Chapter hotspots", f"Top {top_n} factual-drift chapters per model and method.", len(HOTSPOT_HEADERS))
    add_table(hotspots, HOTSPOT_HEADERS, hotspot_data)
    style_sheet(hotspots, len(hotspot_data) + 3, len(HOTSPOT_HEADERS))

    write_title(definitions, "Definitions", "How this workbook calculates each analysis view.", 2)
    add_table(definitions, ["Metric", "Definition"], DEFINITION_ROWS)
    style_sheet(definitions, len(DEFINITION_ROWS) + 3, 2)

    overview.sheet_properties.tabColor = "1F4E78"
    iteration.sheet_properties.tabColor = "5B9BD5"
    criterion.sheet_properties.tabColor = "5B9BD5"
    delta.sheet_properties.tabColor = "70AD47"
    hotspots.sheet_properties.tabColor = "C55A11"
    definitions.sheet_properties.tabColor = "A6A6A6"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def write_csv(path: Path, headers: list[str], rows: list[list[object]]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate deep failure-analysis tables from DysText result logs."
    )
    parser.add_argument(
        "--base-dir",
        type=Path,
        default=Path(__file__).resolve().parents[1],
        help="AIED 2026 directory containing the results folder.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output .xlsx path. Defaults to results/deep_failure_analysis.xlsx.",
    )
    parser.add_argument(
        "--top-n-hotspots",
        type=int,
        default=10,
        help="Number of factual-drift chapter hotspots to keep per model and mode.",
    )
    parser.add_argument(
        "--write-csv",
        action="store_true",
        help="Also write CSV versions of the analysis tables under results/deep_analysis_csv.",
    )
    args = parser.parse_args()

    base_dir = args.base_dir
    output_path = args.output or (base_dir / "results" / "deep_failure_analysis.xlsx")
    rows = annotate_rows(base_dir)
    write_workbook(rows, output_path, args.top_n_hotspots)

    if args.write_csv:
        csv_dir = base_dir / "results" / "deep_analysis_csv"
        write_csv(csv_dir / "overview.csv", OVERVIEW_HEADERS, overview_rows(rows))
        write_csv(csv_dir / "by_iteration.csv", ITERATION_HEADERS, iteration_rows(rows))
        write_csv(csv_dir / "by_failed_check.csv", CRITERION_HEADERS, criterion_rows(rows))
        write_csv(csv_dir / "method_delta.csv", DELTA_HEADERS, method_delta_rows(rows))
        write_csv(csv_dir / "chapter_hotspots.csv", HOTSPOT_HEADERS, hotspot_rows(rows, args.top_n_hotspots))

    print(f"Wrote {output_path}")
