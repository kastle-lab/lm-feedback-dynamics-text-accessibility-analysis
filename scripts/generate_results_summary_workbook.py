# Generate an Excel summary workbook from DysText experiment result logs

from __future__ import annotations

import argparse
import ast
import csv
import re
from collections import Counter
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


RESULT_SETS = [
    {
        "mode": "Holistic",
        "root": "results/local_holistic",
        "pattern": "*_agent_log.csv",
    },
    {
        "mode": "Incremental",
        "root": "results/local_incremental",
        "pattern": "*_agent_log.csv",
    },
    {
        "mode": "Holistic",
        "model": "gpt-global",
        "path": "results/new_agent_log_hollistic.csv",
    },
    {
        "mode": "Incremental",
        "model": "gpt-global",
        "path": "results/all_agent_log_incremental.csv",
    },
]


SUMMARY_HEADERS = [
    "Model",
    "Holistic Pass Rate",
    "Incremental Pass Rate",
    "Holistic Final Pass",
    "Incremental Final Pass",
    "Holistic Avg LIX",
    "Incremental Avg LIX",
    "Holistic Avg DysText",
    "Incremental Avg DysText",
    "Holistic Drift Fails",
    "Incremental Drift Fails",
]


DETAIL_HEADERS = [
    "Mode",
    "Model",
    "Rows",
    "Chapters",
    "Unique Chapter Keys",
    "Checker Passed Iterations",
    "Checker Pass Rate",
    "Final Passed Chapters",
    "Any Passed Chapters",
    "Avg LIX",
    "Avg DysText Visual",
    "Avg DysText Content",
    "Avg DysText Score",
    "Factual Drift Failures",
    "Longphrase Failures",
    "Failed Check Labels",
    "Source Log",
]


METRIC_DEFINITIONS = [
    [
        "Checker Pass Rate",
        "Share of logged iterations where the rewrite passed hard checker rules.",
        "Iteration-level metric. A chapter can pass in one iteration and fail later.",
    ],
    [
        "Final Pass",
        "Number of chapters whose final logged iteration passed the checker.",
        "Chapter-level metric. This is stricter than iteration pass rate.",
    ],
    [
        "Any Passed Chapters",
        "Number of chapters that passed the checker at least once in any iteration.",
        "Useful for seeing whether a model reached a good state before later changes.",
    ],
    [
        "Avg LIX",
        "Average LIX readability score over iterations that reached the judge stage.",
        "Lower values generally mean easier readability.",
    ],
    [
        "Avg DysText Score",
        "Average DysText visual plus content score over iterations that reached the judge stage.",
        "Higher values indicate stronger DysText score under the project scoring function.",
    ],
    [
        "Factual Drift Failures",
        "Count of iterations flagged for changing or losing factual meaning.",
        "Lower is better.",
    ],
    [
        "Holistic",
        "Run mode where all judge feedback is used for the next rewrite.",
        "Applies to local_holistic logs and GPT-global holistic log.",
    ],
    [
        "Incremental",
        "Run mode where only a smaller subset of feedback is used per iteration.",
        "Applies to local_incremental logs and GPT-global incremental log.",
    ],
]


def as_bool(value: object) -> bool:
    return str(value).strip().lower() == "true"


def as_float(value: object) -> float | None:
    if value in ("", None):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def average(values: Iterable[float | None]) -> float | None:
    numbers = [value for value in values if value is not None]
    if not numbers:
        return None
    return sum(numbers) / len(numbers)


def labels(value: object) -> list[str]:
    if not value:
        return []
    try:
        parsed = ast.literal_eval(str(value))
        if isinstance(parsed, list):
            return [str(item) for item in parsed]
    except (SyntaxError, ValueError):
        pass
    return [item.strip() for item in str(value).split(",") if item.strip()]


def dystext_parts(value: object) -> tuple[float | None, float | None, float | None]:
    numbers = re.findall(r"-?\d+(?:\.\d+)?", str(value or ""))
    if len(numbers) < 2:
        return None, None, as_float(value)
    visual = float(numbers[0])
    content = float(numbers[-1])
    return visual, content, visual + content


def model_name_from_log(path: Path) -> str:
    name = path.name
    if name.endswith("_agent_log.csv"):
        name = name[: -len("_agent_log.csv")]
    return name


def read_rows(path: Path) -> list[dict[str, str]]:
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def summarize_log(path: Path, mode: str, model: str | None = None) -> dict[str, object]:
    rows = read_rows(path)
    if not rows:
        raise ValueError(f"No rows found in {path}")

    reached = [row for row in rows if as_bool(row.get("reached_judge"))]
    chapters = [rows[i : i + 4] for i in range(0, len(rows), 4)]
    failed_labels = Counter()
    for row in rows:
        failed_labels.update(labels(row.get("failed_checks")))

    dystext_values = [dystext_parts(row.get("dystext_score")) for row in reached]
    model_name = model or rows[0].get("model") or model_name_from_log(path)

    return {
        "Mode": mode,
        "Model": model_name,
        "Rows": len(rows),
        "Chapters": len(chapters),
        "Unique Chapter Keys": len({row.get("chapter_key") for row in rows if row.get("chapter_key")}),
        "Checker Passed Iterations": len(reached),
        "Checker Pass Rate": len(reached) / len(rows),
        "Final Passed Chapters": sum(
            1 for chapter in chapters if chapter and as_bool(chapter[-1].get("reached_judge"))
        ),
        "Any Passed Chapters": sum(
            1
            for chapter in chapters
            if any(as_bool(row.get("reached_judge")) for row in chapter)
        ),
        "Avg LIX": average(as_float(row.get("lix_score")) for row in reached),
        "Avg DysText Visual": average(value[0] for value in dystext_values),
        "Avg DysText Content": average(value[1] for value in dystext_values),
        "Avg DysText Score": average(value[2] for value in dystext_values),
        "Factual Drift Failures": failed_labels.get("factual_drift", 0),
        "Longphrase Failures": failed_labels.get("longphrase", 0),
        "Failed Check Labels": ";".join(
            f"{label}:{count}" for label, count in sorted(failed_labels.items())
        ),
        "Source Log": str(path),
    }


def collect_metrics(base_dir: Path) -> list[dict[str, object]]:
    metrics: list[dict[str, object]] = []
    for spec in RESULT_SETS:
        mode = spec["mode"]
        if "path" in spec:
            path = base_dir / spec["path"]
            metrics.append(summarize_log(path, mode, spec.get("model")))
            continue

        root = base_dir / spec["root"]
        for path in sorted(root.glob(spec["pattern"])):
            metrics.append(summarize_log(path, mode))
    return metrics


def pivot_summary(metrics: list[dict[str, object]]) -> list[list[object]]:
    by_model: dict[str, dict[str, dict[str, object]]] = {}
    for row in metrics:
        by_model.setdefault(str(row["Model"]), {})[str(row["Mode"])] = row

    rows = []
    for model in sorted(by_model):
        holistic = by_model[model].get("Holistic", {})
        incremental = by_model[model].get("Incremental", {})
        rows.append(
            [
                model,
                holistic.get("Checker Pass Rate"),
                incremental.get("Checker Pass Rate"),
                fraction(holistic.get("Final Passed Chapters"), holistic.get("Chapters")),
                fraction(incremental.get("Final Passed Chapters"), incremental.get("Chapters")),
                holistic.get("Avg LIX"),
                incremental.get("Avg LIX"),
                holistic.get("Avg DysText Score"),
                incremental.get("Avg DysText Score"),
                holistic.get("Factual Drift Failures"),
                incremental.get("Factual Drift Failures"),
            ]
        )
    return rows


def fraction(numerator: object, denominator: object) -> str | None:
    if numerator is None or denominator is None:
        return None
    return f"{numerator}/{denominator}"


def apply_common_styles(ws, rows: int, cols: int, percent_cols: set[int] | None = None):
    percent_cols = percent_cols or set()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin_gray)

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

    for cell in ws[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=4, max_row=rows, max_col=cols):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if cell.column in percent_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%"
            elif isinstance(cell.value, float):
                cell.number_format = "0.00"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"

    for col_idx in range(1, cols + 1):
        letter = get_column_letter(col_idx)
        max_len = max(
            len(str(ws.cell(row=row, column=col_idx).value or ""))
            for row in range(1, rows + 1)
        )
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 34)


def write_title(ws, title: str, subtitle: str, cols: int):
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True, color="1F1F1F")
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="666666")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(cols, 6))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(cols, 8))


def add_table(ws, headers: list[str], rows: list[list[object]], start_row: int = 3):
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=start_row, column=col_idx, value=header)
    for row_idx, row in enumerate(rows, start=start_row + 1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def add_chart(ws, rows: int):
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Checker pass rate by model"
    chart.y_axis.title = "Model"
    chart.x_axis.title = "Pass rate"
    data = Reference(ws, min_col=2, max_col=3, min_row=3, max_row=rows)
    cats = Reference(ws, min_col=1, min_row=4, max_row=rows)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 9
    chart.width = 18
    ws.add_chart(chart, "M3")


def create_workbook(metrics: list[dict[str, object]], output_path: Path):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    detail = wb.create_sheet("Metrics")
    definitions = wb.create_sheet("Definitions")

    summary_rows = pivot_summary(metrics)
    write_title(summary, "Model-wise results summary", "Generated from local result logs.", len(SUMMARY_HEADERS))
    add_table(summary, SUMMARY_HEADERS, summary_rows)
    apply_common_styles(summary, len(summary_rows) + 3, len(SUMMARY_HEADERS), {2, 3})
    add_chart(summary, len(summary_rows) + 3)

    detail_rows = [[row.get(header) for header in DETAIL_HEADERS] for row in sorted(metrics, key=lambda r: (str(r["Model"]), str(r["Mode"])))]
    write_title(detail, "Detailed metrics", "One row per model and feedback mode.", len(DETAIL_HEADERS))
    add_table(detail, DETAIL_HEADERS, detail_rows)
    apply_common_styles(detail, len(detail_rows) + 3, len(DETAIL_HEADERS), {7})
    detail.freeze_panes = "A4"

    write_title(definitions, "Metric definitions", "Definitions used by the analysis script.", 3)
    add_table(definitions, ["Metric", "Meaning", "Notes"], METRIC_DEFINITIONS)
    apply_common_styles(definitions, len(METRIC_DEFINITIONS) + 3, 3)
    definitions.freeze_panes = "A4"

    summary.sheet_properties.tabColor = "1F4E78"
    detail.sheet_properties.tabColor = "5B9BD5"
    definitions.sheet_properties.tabColor = "A6A6A6"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate a model-wise XLSX summary from DysText result logs."
    )
    parser.add_argument(
        "--base-dir",
        type=Path,
        default=Path(__file__).resolve().parents[1],
        help="AIED 2026 directory containing the results folder.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output .xlsx path. Defaults to results/model_wise_summary.xlsx.",
    )
    args = parser.parse_args()

    base_dir = args.base_dir
    output_path = args.output or (base_dir / "results" / "model_wise_summary.xlsx")
    metrics = collect_metrics(base_dir)
    create_workbook(metrics, output_path)
    print(f"Wrote {output_path}")
